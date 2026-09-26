from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


def add_sheet_from_rows(wb: Workbook, title: str, rows: list[dict[str, Any]]):
    ws = wb.create_sheet(title[:31])
    if not rows:
        ws.append(["No data"])
        return ws

    headers = []
    for row in rows:
        for key in row:
            if key not in headers:
                headers.append(key)

    ws.append(headers)
    for row in rows:
        ws.append([row.get(header, "") for header in headers])

    return ws


def format_workbook(wb: Workbook):
    header_fill = PatternFill(fill_type="solid", fgColor="0F766E")
    header_font = Font(bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_alignment = Alignment(vertical="top", wrap_text=True)
    border = Border(bottom=Side(style="thin", color="D1D5DB"))

    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = border
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = body_alignment
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 32

        for column_cells in ws.columns:
            letter = column_cells[0].column_letter
            max_len = 0
            for cell in column_cells[:80]:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
            ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 45)

        if ws.max_row >= 2 and ws.max_column >= 1:
            ref = ws.dimensions
            safe_name = "Table" + str(len(wb.worksheets)) + "_" + str(ws._current_row)
            safe_name = "T" + "".join(ch for ch in safe_name if ch.isalnum())[:20]
            table = Table(displayName=safe_name, ref=ref)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium4",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            ws.add_table(table)


def export_case(case_id: str, title: str, data: dict[str, Any], output_path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    rows = data.get("results", [])
    findings = data.get("findings") or data.get("mixed") or []

    ws.append(["Test Case ID", "Title", "Pages/Data Rows", "Findings", "Overall Status"])
    overall_status = "PASS"
    if any(str(row.get("Status", "")).upper() == "FAIL" for row in rows):
        overall_status = "FAIL"
    elif findings:
        overall_status = "FAIL"
    ws.append([case_id, title, len(rows), len(findings), overall_status])

    add_sheet_from_rows(wb, "Results", rows)
    if findings:
        add_sheet_from_rows(wb, "Findings", findings)

    format_workbook(wb)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"Excel report: {output_path.resolve()}")
