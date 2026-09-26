from __future__ import annotations

import argparse
import importlib
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from common import REPORT_DIR

TESTS = [
    ("TC-SEC-01", "HTTPS Enforcement and Mixed Content", "tests.tc_sec_01_https_mixed_content"),
    ("TC-SEC-02", "Security Response Headers", "tests.tc_sec_02_security_headers"),
    ("TC-SEC-03", "Public Input Validation and Sanitization", "tests.tc_sec_03_input_validation"),
    ("TC-SEC-04", "Sensitive Information Exposure", "tests.tc_sec_04_sensitive_exposure"),
    ("TC-SEC-05", "Cookie Security Attributes", "tests.tc_sec_05_cookie_security"),
]


def build_combined(results):
    output = REPORT_DIR / "Thaura_Security_Testing_Combined_Report.xlsx"
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"

    summary_headers = [
        "Test Case ID",
        "Title",
        "Status",
        "Rows Collected",
        "Findings",
        "Excel Detail File",
    ]
    summary.append(summary_headers)

    for item in results:
        summary.append([
            item["id"],
            item["title"],
            item["status"],
            item["rows"],
            item["findings"],
            item["detail_file"],
        ])

    for test_id, title, module_name in TESTS:
        item = next(x for x in results if x["id"] == test_id)
        data = item["data"]

        rows = data.get("results", [])
        findings = data.get("findings") or data.get("mixed") or []

        ws = wb.create_sheet(test_id)
        if rows:
            headers = []
            for row in rows:
                for key in row:
                    if key not in headers:
                        headers.append(key)
            ws.append(headers)
            for row in rows:
                ws.append([row.get(h, "") for h in headers])
        else:
            ws.append(["No data"])

        if findings:
            fw = wb.create_sheet((test_id + " Findings")[:31])
            fheaders = []
            for row in findings:
                for key in row:
                    if key not in fheaders:
                        fheaders.append(key)
            fw.append(fheaders)
            for row in findings:
                fw.append([row.get(h, "") for h in fheaders])

    header_fill = PatternFill(fill_type="solid", fgColor="0F766E")
    header_font = Font(bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_alignment = Alignment(vertical="top", wrap_text=True)
    border = Border(bottom=Side(style="thin", color="D1D5DB"))

    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = border
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = body_alignment
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            letter = col[0].column_letter
            max_len = max((len(str(c.value or "")) for c in col[:80]), default=10)
            ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 45)
        if ws.max_row >= 2:
            table_name = "T" + "".join(ch for ch in ws.title if ch.isalnum())[:20]
            ws.add_table(Table(
                displayName=table_name,
                ref=ws.dimensions,
                tableStyleInfo=TableStyleInfo(
                    name="TableStyleMedium4",
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False,
                )
            ))

    wb.save(output)
    return output


def main(show_browser: bool):
    results = []

    for test_id, title, module_name in TESTS:
        print("\n" + "=" * 110)
        print(f"RUNNING {test_id}: {title}")
        print("=" * 110)

        module = importlib.import_module(module_name)
        data = module.run(show_browser=show_browser)

        rows = data.get("results", [])
        findings = data.get("findings") or data.get("mixed") or []

        status = "PASS"
        if findings or any(str(row.get("Status", "")).upper() == "FAIL" for row in rows):
            status = "FAIL"

        detail_file = REPORT_DIR / f"{test_id}_detail.xlsx"

        # Write a compact per-case workbook using the same collected data.
        from excel_utils import export_case
        export_case(
            test_id,
            title,
            data,
            detail_file,
        )

        results.append({
            "id": test_id,
            "title": title,
            "status": status,
            "rows": len(rows),
            "findings": len(findings),
            "detail_file": str(detail_file.resolve()),
            "data": data,
        })

    combined = build_combined(results)

    print("\n" + "=" * 110)
    print("SECURITY AUTOMATION COMPLETED")
    print("=" * 110)
    for item in results:
        print(
            f"{item['id']} | {item['status']} | rows={item['rows']} | findings={item['findings']}"
        )
    print(f"Combined Excel: {combined.resolve()}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Run Thaura basic non-intrusive security automation"
    )
    parser.add_argument("-sv", "--show-browser", action="store_true")
    args = parser.parse_args()
    main(args.show_browser)
